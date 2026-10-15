"""excel_sync.py（CSV書き出し）と excel_check.py（読み取り専用の照合）のテスト。ダミーデータのみ。
python -m unittest test_excel_sync -v"""
import csv
import os
import tempfile
import unittest

import openpyxl
import pandas as pd

import excel_check
import excel_sync
from test_record_edit import BASE, make_df


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self._saved = (excel_sync.EXCEL_PATH, excel_sync.EXPORT_DIR)
        excel_sync.EXCEL_PATH = os.path.join(self.tmp, "dummy.xlsm")
        excel_sync.EXPORT_DIR = self.tmp

    def tearDown(self):
        excel_sync.EXCEL_PATH, excel_sync.EXPORT_DIR = self._saved

    def test_csv_content_and_encoding(self):
        df = make_df(BASE)
        df.loc[df["id"] == 2, "note"] = 'カンマ,と"引用符"\n改行'
        df.loc[df["id"] == 3, "fuel_unit_price"] = float("nan")
        res = excel_sync.export_records_for_excel(df)
        self.assertEqual(res["status"], "exported")
        self.assertEqual(res["count"], 4)
        with open(res["path"], "rb") as f:
            self.assertTrue(f.read(3) == b"\xef\xbb\xbf")          # UTF-8 BOM
        with open(res["path"], encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
        self.assertEqual([r["id"] for r in rows], ["1", "2", "3", "4"])
        self.assertEqual(rows[0]["record_date"], "2026-09-01")
        self.assertEqual(rows[0]["odometer_km"], "10000")
        self.assertEqual(rows[1]["note"], 'カンマ,と"引用符"\n改行')
        self.assertEqual(rows[2]["fuel_unit_price"], "")
        self.assertTrue(os.path.exists(os.path.join(self.tmp, "sync_status.txt")))

    def test_not_configured_when_no_local_settings(self):
        excel_sync.EXCEL_PATH = None
        self.assertEqual(excel_sync.export_records_for_excel(make_df(BASE))["status"], "not_configured")

    def test_rows_sorted_by_date_then_id(self):
        df = make_df(BASE).iloc[::-1].reset_index(drop=True)
        res = excel_sync.export_records_for_excel(df)
        with open(res["path"], encoding="utf-8-sig", newline="") as f:
            ids = [r["id"] for r in csv.DictReader(f)]
        self.assertEqual(ids, ["1", "2", "3", "4"])


class CheckTests(unittest.TestCase):
    def make_book(self, rows):
        path = os.path.join(tempfile.mkdtemp(), "dummy.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "燃費確認表"
        for i, (rid, d, fuel, odo, price) in enumerate(rows, start=2):
            ws.cell(i, 1, rid)
            ws.cell(i, 2, d)
            ws.cell(i, 3, fuel)
            ws.cell(i, 4, odo)
            if price is not None:
                ws.cell(i, 6, price)
        wb.save(path)
        return path

    def test_all_match(self):
        from datetime import datetime
        df = make_df(BASE)
        rows = [(int(r.id), r.record_date.to_pydatetime(), r.fuel_liters, r.odometer_km, r.fuel_unit_price)
                for r in df.itertuples()]
        res = excel_check.compare(df, self.make_book(rows))
        self.assertIn("一致 4 件", res["summary"])
        self.assertEqual(res["rows"], [])

    def test_differences_are_listed(self):
        df = make_df(BASE)
        rows = [(int(r.id), r.record_date.to_pydatetime(), r.fuel_liters, r.odometer_km, r.fuel_unit_price)
                for r in df.itertuples()]
        rows[1] = (2, rows[1][1], rows[1][2], 10450.0, rows[1][4])          # 値が違う
        rows.pop()                                                          # クラウドだけ（ID4）
        rows.append((99, rows[0][1], 30.0, 9999.0, 150.0))                  # Excelだけ
        res = excel_check.compare(df, self.make_book(rows))
        kinds = sorted(r["区分"] for r in res["rows"])
        self.assertEqual(kinds, ["Excelだけにある", "クラウドだけにある", "値が違う"])

    def test_missing_ids_message(self):
        df = make_df(BASE)
        rows = [(None, r.record_date.to_pydatetime(), r.fuel_liters, r.odometer_km, r.fuel_unit_price)
                for r in df.itertuples()]
        res = excel_check.compare(df, self.make_book(rows))
        self.assertIn("IDがありません", res["summary"])


if __name__ == "__main__":
    unittest.main()
