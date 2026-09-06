"""
既存の「ソリオ燃費早見表.xlsm」から燃費管理アプリ(fuel_log.db)へ
過去の記録を一括インポートするスクリプト。

使い方:
    python import_from_excel.py "対象のxlsmファイルパス"

- シート「燃費確認表」のB〜I列（日付・給油量・メーター値・燃費・単価・金額・メモ・区間距離）を読み取ります。
- 既にfuel_log.dbに同じ日付の記録がある場合はスキップします（重複インポート防止）。
- 実行後、streamlit run app.py で開くと「履歴・グラフ」タブに反映されます。
"""

import os
import sqlite3
import sys
from datetime import datetime, timedelta

import openpyxl

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(BASE_DIR, "fuel_log.db")

SHEET_NAME = "燃費確認表"
COL_DATE = 2       # B: 日付
COL_FUEL = 3        # C: 給油量
COL_ODO = 4          # D: 走行距離メーター値
COL_EFF = 5          # E: 燃費(km/L)
COL_PRICE = 6        # F: 給油単価(円)
COL_COST = 7         # G: 給油金額(円)
COL_NOTE = 8         # H: オイル交換時期/メモ
COL_DIST = 9         # I: 区間走行距離(自動計算値)


def to_date(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, (int, float)):
        return (datetime(1899, 12, 30) + timedelta(days=v)).date()
    return None


def to_number(v):
    if isinstance(v, (int, float)):
        return float(v)
    return None


def init_db():
    conn = sqlite3.connect(DB_PATH)
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS records (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            record_date TEXT NOT NULL,
            odometer_km REAL NOT NULL,
            distance_km REAL,
            fuel_liters REAL NOT NULL,
            fuel_unit_price REAL,
            fuel_cost REAL,
            efficiency_km_per_l REAL,
            odometer_photo TEXT,
            receipt_photo TEXT,
            note TEXT,
            created_at TEXT NOT NULL
        )
        """
    )
    conn.commit()
    return conn


def existing_dates(conn):
    return {row[0] for row in conn.execute("SELECT record_date FROM records").fetchall()}


def main():
    if len(sys.argv) < 2:
        print("使い方: python import_from_excel.py <xlsmファイルパス>")
        sys.exit(1)

    xlsm_path = sys.argv[1]
    wb = openpyxl.load_workbook(xlsm_path, data_only=True, keep_vba=True)
    if SHEET_NAME not in wb.sheetnames:
        print(f"シート「{SHEET_NAME}」が見つかりません。シート一覧: {wb.sheetnames}")
        sys.exit(1)
    ws = wb[SHEET_NAME]

    conn = init_db()
    already = existing_dates(conn)

    inserted, skipped = 0, 0
    prev_odo = None

    for r in range(2, ws.max_row + 1):
        raw_date = ws.cell(row=r, column=COL_DATE).value
        d = to_date(raw_date)
        if d is None:
            continue

        fuel = to_number(ws.cell(row=r, column=COL_FUEL).value)
        odo = to_number(ws.cell(row=r, column=COL_ODO).value)
        if fuel is None or odo is None:
            continue

        eff = to_number(ws.cell(row=r, column=COL_EFF).value)
        price = to_number(ws.cell(row=r, column=COL_PRICE).value)
        cost = to_number(ws.cell(row=r, column=COL_COST).value)
        note_val = ws.cell(row=r, column=COL_NOTE).value
        note = note_val if isinstance(note_val, str) and note_val.strip() else None

        dist_val = to_number(ws.cell(row=r, column=COL_DIST).value)
        if dist_val is None and prev_odo is not None:
            dist_val = round(odo - prev_odo, 1)
        prev_odo = odo

        iso_date = d.isoformat()
        if iso_date in already:
            skipped += 1
            continue

        conn.execute(
            """
            INSERT INTO records
            (record_date, odometer_km, distance_km, fuel_liters, fuel_unit_price, fuel_cost,
             efficiency_km_per_l, odometer_photo, receipt_photo, note, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, NULL, NULL, ?, ?)
            """,
            (
                iso_date, odo, dist_val, fuel, price, cost, eff, note,
                datetime.now().isoformat(timespec="seconds"),
            ),
        )
        already.add(iso_date)
        inserted += 1

    conn.commit()
    conn.close()
    print(f"インポート完了: {inserted}件追加, {skipped}件は既存のためスキップ")


if __name__ == "__main__":
    main()
